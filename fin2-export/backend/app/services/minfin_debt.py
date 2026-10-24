"""Госдолг субъектов РФ и его структура — из помесячных таблиц Минфина России.

Минфин публикует по каждому субъекту объём долга и его состав: государственные ценные
бумаги, кредиты кредитных организаций, бюджетные кредиты из других бюджетов, госгарантии
и иные обязательства (https://minfin.gov.ru/ru/perfomance/public_debt/subj/). Это первичный
источник — он точнее, чем пересказ в пресс-релизах агентств, и покрывает все регионы сразу.

Файлы лежат по нерегулярным адресам, поэтому ссылки на нужные отчётные даты держим списком.
Суммы в таблицах — тысячи рублей, мы приводим к миллионам, как остальные данные сектора.
Результат кэшируется в services/data/subfed_debt.json; обновлять раз в квартал скриптом
scripts/subfed_fill.py.
"""
from __future__ import annotations

import json
import logging
import os
import re
import warnings

import requests

log = logging.getLogger("minfin_debt")
warnings.filterwarnings("ignore", message="Unverified HTTPS request")

_CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "subfed_debt.json")
_UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 "
       "(KHTML, like Gecko) Version/17.0 Safari/605.1.15")
_BASE = "https://minfin.gov.ru"

# отчётная дата → файл Минфина. Период в наших данных: год = состояние на 1 января
# следующего года, полугодие — на 1 июля.
FILES: dict[str, str] = {
    "2021": "/common/upload/library/2022/02/main/DO_na_01.01.2022.xlsx",
    "2022": "/common/upload/library/2023/01/main/DO_na_01.01.2023.xlsx",
    "2023": "/common/upload/library/2024/07/main/DO_na_01.01.2024_ot_18.07.2024.xlsx",
    "2024": "/common/upload/library/2025/04/main/Obem_i_struktura_na_1_yanvarya_2025_g.xlsx",
    "2025": "/common/upload/library/2026/04/main/DO_na_01.01.2026_ot_20.04.2026.xlsx",
    "6м2026": "/common/upload/library/2026/07/main/DO_na_01.07.2026.xlsx",
}

# строки-итоги федеральных округов и России в таблице пропускаем
_SKIP = re.compile(r"(федеральн\w+ округ|итого|российск\w+ федерац)", re.I)


def _norm(name: str) -> str:
    """Имя субъекта к общему виду: у Минфина «г.Москва», «Ханты-Мансийский АО - Югра»."""
    x = str(name or "").strip().lower().replace("ё", "е")
    x = re.sub(r"[«»\"()]", " ", x)
    x = re.sub(r"[—–\-]", " ", x)
    x = re.sub(r"\b(г|гор|город)\.?\b", " ", x)
    x = re.sub(r"\bавтономн(ый|ая)\b", " ", x)
    x = re.sub(r"\b(округ|область|край|республика|республики)\b", " ", x)
    x = x.replace("а о", " ").replace(" ао ", " ")
    w = [t for t in x.split() if t and t not in ("ао",)]
    return " ".join(t for i, t in enumerate(w) if i == 0 or t != w[i - 1])


def _rows(content: bytes) -> dict[str, dict]:
    import io as _io

    import openpyxl
    wb = openpyxl.load_workbook(_io.BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(values_only=True):
        name = row[0]
        if not isinstance(name, str) or not name.strip() or _SKIP.search(name):
            continue
        nums = [x if isinstance(x, (int, float)) else None for x in row[1:7]]
        if nums[0] is None and nums[2] is None:          # строка-заголовок
            continue
        mln = lambda v: None if v is None else round(v / 1000, 1)   # тыс. ₽ → млн ₽  # noqa: E731
        bonds, bank, budget, guar, other, total = (nums + [None] * 6)[:6]
        out[_norm(name)] = {"name": name.strip(), "bonds": mln(bonds), "bank_loans": mln(bank),
                            "budget_loans": mln(budget), "guarantees": mln(guar),
                            "other": mln(other), "debt": mln(total)}
    return out


def fetch(periods: list[str] | None = None) -> dict:
    """Скачать таблицы Минфина по нужным отчётным датам и положить в кэш."""
    out = {"source": _BASE + "/ru/perfomance/public_debt/subj/", "units": "млн ₽", "periods": {}}
    for period, path in FILES.items():
        if periods and period not in periods:
            continue
        r = requests.get(_BASE + path, headers={"User-Agent": _UA}, timeout=90, verify=False)
        r.raise_for_status()
        rows = _rows(r.content)
        out["periods"][period] = rows
        log.info("Минфин: %s — %s субъектов", period, len(rows))
    try:
        with open(_CACHE, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False)
    except OSError as e:  # noqa: BLE001
        log.warning("кэш долга субъектов: %s", e)
    return out


def load() -> dict:
    try:
        with open(_CACHE, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {"periods": {}}


def for_region(name: str) -> dict:
    """{период: {долг и его структура}} по имени субъекта в наших данных."""
    key = _norm(name)
    data = load().get("periods") or {}
    return {p: rows[key] for p, rows in data.items() if key in rows}
