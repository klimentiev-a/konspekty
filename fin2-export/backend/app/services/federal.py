"""Российская Федерация как эмитент: федеральный бюджет и государственный внутренний долг.

Два первоисточника Минфина России:
  • «Краткая информация об исполнении федерального бюджета» (приложение 3, квартальные
    данные накопленным итогом) — доходы с разбивкой на нефтегазовые и ненефтегазовые,
    расходы, обслуживание долга и дефицит. Суммы в млрд ₽;
  • «Объем и структура государственного внутреннего долга» (debtDynamics) — остатки по видам
    ОФЗ: ПД (постоянный купон), ПК (плавающий), ИН (индексируемый номинал), АД (амортизация)
    и ГСО. Тоже млрд ₽, помесячно.

Рыночную часть — объём ОФЗ в обращении, дюрацию и кривую — берём из своих данных, как по
остальным эмитентам. Кэш: services/data/federal.json, обновляется скриптом subfed_fill.py.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import logging
import os
import re
import warnings

import requests

log = logging.getLogger("federal")
warnings.filterwarnings("ignore", message="Unverified HTTPS request")

_CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "federal.json")
_UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 "
       "(KHTML, like Gecko) Version/17.0 Safari/605.1.15")
_BUDGET = "https://minfin.gov.ru/ru/statistics/fedbud/"
_DEBT = "https://minfin.gov.ru/ru/perfomance/public_debt/internal/structure/"

# строки приложения 3 → наши ключи
_ROWS = {"1": "inc", "1.1.": "oil", "1.2.": "nonoil", "2": "exp", "2.13.": "svc", "3": "bal"}
# наш период → колонка квартального файла (накопленным итогом)
_QCOL = {"2021": "4кв.2021", "2022": "4кв.2022", "2023": "4кв.2023", "2024": "4кв.2024",
         "2025": "4кв.2025", "6м2026": "2 кв.2026"}


def _find(page: str, pattern: str) -> str | None:
    r = requests.get(page, headers={"User-Agent": _UA}, timeout=40, verify=False)
    r.raise_for_status()
    m = re.search(pattern, r.text)
    return "https://minfin.gov.ru" + m.group(1) if m else None


def _budget() -> dict:
    url = _find(_BUDGET, r'href="([^"]*Prilozhenie_3[^"]*kvartal[^"]*\.xlsx)"')
    if not url:
        raise RuntimeError("не найден квартальный файл исполнения федерального бюджета")
    import openpyxl
    blob = requests.get(url, headers={"User-Agent": _UA}, timeout=120, verify=False).content
    ws = openpyxl.load_workbook(io.BytesIO(blob), data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[2]
    # «2 кв.2026 **» — сноска в шапке, сравниваем по началу строки
    cols = {str(v).replace(" ", "").split("*")[0]: i for i, v in enumerate(hdr) if v}
    out: dict[str, dict] = {}
    for row in rows:
        key = _ROWS.get(str(row[0] or "").strip())
        if not key:
            continue
        for period, col in _QCOL.items():
            i = cols.get(col.replace(" ", ""))
            if i is None:
                continue
            v = row[i]
            if isinstance(v, (int, float)):
                out.setdefault(period, {})[key] = round(v * 1000, 1)     # млрд ₽ → млн ₽
    return out


def _debt_history() -> dict:
    """Длинный ряд структуры внутреннего долга по годам (с 1993-го, нам нужен хвост)."""
    url = _find(_DEBT, r'href="([^"]*Obem_gos_vnutr__dolga_s_93_goda[^"]*\.xlsx)"')
    if not url:
        return {}
    import openpyxl
    blob = requests.get(url, headers={"User-Agent": _UA}, timeout=120, verify=False).content
    ws = openpyxl.load_workbook(io.BytesIO(blob), data_only=True).worksheets[0]
    head, idx, out = None, {}, {}
    for row in ws.iter_rows(values_only=True):
        cells = [str(x).strip() if x is not None else "" for x in row]
        if head is None and any(c == "ОФЗ-ПД" for c in cells):
            head, idx = cells, {c: i for i, c in enumerate(cells) if c}
            continue
        if head is None:
            continue
        m = re.match(r"на 01\.01\.(\d{4})", cells[1] if len(cells) > 1 else "")
        if not m:
            continue
        period = str(int(m.group(1)) - 1)          # состояние на 1 января → прошлый год
        num = lambda k: (row[idx[k]] if k in idx and isinstance(row[idx[k]], (int, float)) else None)  # noqa: E731
        vals = {"ofz_pd": num("ОФЗ-ПД"), "ofz_pk": num("ОФЗ-ПК"), "ofz_in": num("ОФЗ-ИН"),
                "ofz_ad": num("ОФЗ-АД"), "gso": num("ГСО-ППС") or num("ГСО-ПДС")}
        vals = {k: (None if v is None else round(v * 1000, 1)) for k, v in vals.items()}
        if any(v is not None for v in vals.values()):
            vals["internal"] = round(sum(v for v in vals.values() if v is not None), 1)
            out[period] = vals
    return out


def _debt() -> dict:
    url = _find(_DEBT, r'href="([^"]*debtDynamics[^"]*\.xlsx)"')
    if not url:
        raise RuntimeError("не найден файл структуры внутреннего долга")
    import openpyxl
    blob = requests.get(url, headers={"User-Agent": _UA}, timeout=120, verify=False).content
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    out: dict[str, dict] = {}
    for ws in wb.worksheets:
        head, idx = None, {}
        for row in ws.iter_rows(values_only=True):
            cells = [str(x) if x is not None else "" for x in row]
            if head is None and any("ОФЗ-ПД" in c for c in cells):
                head = cells
                idx = {c.strip(): i for i, c in enumerate(cells) if c.strip()}
                continue
            if head is None:
                continue
            date = next((x for x in row if isinstance(x, dt.datetime)), None)
            if not date:
                continue
            # год = состояние на 1 января следующего года, полугодие — на 1 июля
            period = (str(date.year - 1) if (date.month, date.day) == (1, 1)
                      else f"6м{date.year}" if (date.month, date.day) == (7, 1) else None)
            if not period:
                continue
            get = lambda k: (row[idx[k]] if k in idx and isinstance(row[idx[k]], (int, float)) else None)  # noqa: E731
            vals = {"ofz_pd": get("ОФЗ-ПД"), "ofz_pk": get("ОФЗ-ПК"), "ofz_in": get("ОФЗ-ИН"),
                    "ofz_ad": get("ОФЗ-АД"), "gso": get("ГСО-ППС"),
                    "internal": next((row[i] for k, i in idx.items() if "Итого" in k
                                      and isinstance(row[i], (int, float))), None)}
            vals = {k: (None if v is None else round(v * 1000, 1)) for k, v in vals.items()}  # млрд → млн
            if any(v is not None for v in vals.values()):
                out[period] = vals
    return out


def fetch() -> dict:
    debt = _debt_history()
    debt.update(_debt())                 # свежие даты из помесячного файла точнее
    out = {"source_budget": _BUDGET, "source_debt": _DEBT, "units": "млн ₽",
           "budget": _budget(), "debt": debt,
           "at": dt.datetime.utcnow().isoformat(timespec="seconds")}
    try:
        with open(_CACHE, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False)
    except OSError as e:  # noqa: BLE001
        log.warning("кэш федеральных данных: %s", e)
    return out


def load() -> dict:
    try:
        with open(_CACHE, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {"budget": {}, "debt": {}}


def by_period() -> dict:
    """{период: {бюджет и долг РФ}} — в том же виде, что показатели регионов."""
    d = load()
    out: dict[str, dict] = {}
    for period in set(list((d.get("budget") or {}).keys()) + list((d.get("debt") or {}).keys())):
        rec = dict((d.get("budget") or {}).get(period) or {})
        rec.update((d.get("debt") or {}).get(period) or {})
        inc, exp, svc = rec.get("inc"), rec.get("exp"), rec.get("svc")
        rec["bal_inc"] = round(100 * rec["bal"] / inc, 1) if (rec.get("bal") is not None and inc) else None
        rec["svc_exp"] = round(100 * svc / exp, 2) if (svc is not None and exp) else None
        rec["oil_share"] = round(100 * rec["oil"] / inc, 1) if (rec.get("oil") is not None and inc) else None
        ofz = [rec.get(k) for k in ("ofz_pd", "ofz_pk", "ofz_in", "ofz_ad", "gso")]
        tot = rec.get("internal") or (sum(x for x in ofz if x is not None) or None)
        rec["internal"] = tot
        rec["pk_share"] = round(100 * rec["ofz_pk"] / tot, 1) if (rec.get("ofz_pk") is not None and tot) else None
        rec["debt_inc"] = round(tot / inc, 3) if (tot and inc) else None
        out[period] = rec
    return out
