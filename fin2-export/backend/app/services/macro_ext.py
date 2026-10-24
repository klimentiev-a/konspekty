"""Макро, вторая волна: реальный сектор Росстата + деньги/ликвидность/внешний сектор ЦБ.

Каждая функция — один источник; сбой одного не роняет остальные (get_macro ловит
исключения по каждому). Имена файлов Росстата и ЦБ меняются каждый месяц
(ipc_mes_08-2026 → ipc_mes_09-2026), поэтому ссылку ищем на странице раздела по шаблону.
У rosstat.gov.ru невалидный SSL → качаем через macro._rosstat_get (verify=False)."""
import io
import os
import re
import threading
import time
import warnings
from datetime import datetime, timedelta

from .macro import _UA, _ds_multi, _dump_json, _eom, _get, _hd_rows, _num, _pick, _rosstat_get, log  # noqa: F401

_RS = "https://rosstat.gov.ru"
_CBR = "https://www.cbr.ru"

# ─── общие помощники ────────────────────────────────────────────────────────────────
_pages: dict = {}
_plock = threading.Lock()


def _page(url: str, rosstat: bool = True) -> str:
    """HTML страницы раздела (кэш 30 мин — одну страницу читают несколько источников)."""
    with _plock:
        hit = _pages.get(url)
        if hit and time.time() - hit[0] < 1800:
            return hit[1]
    txt = _rosstat_get(url).decode("utf-8", "ignore") if rosstat else _get(url).text
    with _plock:
        _pages[url] = (time.time(), txt)
    return txt


def _link(page_url: str, pattern: str, rosstat: bool = True) -> str:
    """Первая ссылка на файл, чьё имя подходит под regex `pattern`."""
    m = re.findall(r'href="([^"]*?(?:' + pattern + r'))"', _page(page_url, rosstat))
    if not m:
        raise RuntimeError(f"нет файла {pattern} на {page_url}")
    u = m[0]
    return u if u.startswith("http") else (_RS if rosstat else _CBR) + u


def _book(content: bytes, xls: bool = False, only=None) -> dict:
    """{лист: [[ячейки]]} для xlsx (openpyxl) или старого xls (xlrd)."""
    if xls:
        import xlrd
        wb = xlrd.open_workbook(file_contents=content)
        return {s.name: [s.row_values(i) for i in range(s.nrows)]
                for s in wb.sheets() if only is None or s.name in only}
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return {n: [list(r) for r in wb[n].iter_rows(values_only=True)]
                for n in wb.sheetnames if only is None or n in only}


def _rs_book(page: str, pattern: str, only=None) -> dict:
    url = _link(page, pattern)
    return _book(_rosstat_get(url), url.endswith(".xls"), only)


def _cbr_book(url: str, only=None) -> dict:
    return _book(_get(url).content, url.endswith(".xls"), only)


def _f(x):
    """Число из ячейки: 101.3 / '101,3' / '62 040' (пробел-разделитель тысяч) / '98,6 2)' (сноска)
    / '-' → None. Без разбора разделителя тысяч «62 040» читалось как 62 (зарплаты 10-12.2022)."""
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace("\xa0", " ").replace(" ", " ").strip()
    s = re.sub(r"\s*(?:\(\d+\)|\d\)|\*+)$", "", s)          # сноски: «62 470(2)», «98,6 2)», «5*»
    if re.fullmatch(r"-?\d{1,3}(?: \d{3})+(?:[.,]\d+)?", s):   # «62 040», «1 234 567,8»
        return float(s.replace(" ", "").replace(",", "."))
    m = re.match(r"\s*(-?\d+(?:[.,]\d+)?)", s)
    return float(m.group(1).replace(",", ".")) if m else None


def _year(x):
    """Год из ячейки заголовка: 2011 / 2011.0 / '2014г.' / '20222)' (сноска) → int."""
    if isinstance(x, datetime) or isinstance(x, bool) or x is None:
        return None
    if isinstance(x, (int, float)):
        return int(x) if 1990 <= x <= 2035 and float(x).is_integer() else None
    m = re.match(r"\s*((?:19|20)\d{2})(?!\d)", str(x))
    if not m:
        m = re.match(r"\s*((?:19|20)\d{2})\d\)", str(x))  # '20222)' — год + сноска
    return int(m.group(1)) if m else None


_MFULL = {m: i for i, m in enumerate(["январь", "февраль", "март", "апрель", "май", "июнь", "июль",
                                        "август", "сентябрь", "октябрь", "ноябрь", "декабрь"], 1)}
_MABBR = {m: i for i, m in enumerate(["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг",
                                        "сен", "окт", "ноя", "дек"], 1)}


def _mon(x):
    """Номер месяца из 'Январь' / ' январь2)' / 'Янв'. Накопленные 'январь-март' → None."""
    if not isinstance(x, str):
        return None
    s = re.sub(r"\d+\)?$", "", x.strip().lower().rstrip(".")).strip()
    return _MFULL.get(s) or (_MABBR.get(s) if len(s) <= 4 else None)


def _ser(d: dict, frm: str = "") -> list:
    return [{"date": k, "val": d[k]} for k in sorted(d) if k >= frm and d[k] is not None]


def _ym_grid(rows: list) -> dict:
    """Годы по колонкам, месяцы по строкам → {'YYYY-MM': v}. Берём первый блок
    (второй блок того же листа, напр. «к декабрю», начинается повтором месяца)."""
    hi = next((i for i, r in enumerate(rows) if sum(1 for c in r[1:] if _year(c)) >= 5), None)
    if hi is None:
        return {}
    ycol = {j: _year(c) for j, c in enumerate(rows[hi]) if j > 0 and _year(c)}
    out, seen = {}, set()
    for r in rows[hi + 1:]:
        m = _mon(r[0]) if r else None
        if not m:
            continue
        if m in seen:
            break
        seen.add(m)
        for j, y in ycol.items():
            v = _f(r[j]) if j < len(r) else None
            if v is not None:
                out[f"{y}-{m:02d}"] = v
    return out


def _yoy_chain(mm: dict) -> dict:
    """Индексы м/м (%, 100=без изменений) → г/г (%) произведением 12 подряд идущих месяцев."""
    ks, out = sorted(mm), {}
    for i in range(11, len(ks)):
        y, m = map(int, ks[i].split("-"))
        y0, m0 = map(int, ks[i - 11].split("-"))
        if (y - y0) * 12 + (m - m0) != 11:
            continue
        p = 1.0
        for k in ks[i - 11:i + 1]:
            p *= mm[k] / 100
        out[ks[i]] = round((p - 1) * 100, 2)
    return out


def _mom(mm: dict) -> dict:
    return {k: round(v - 100, 2) for k, v in mm.items()}


def _prev_year(k: str) -> str:
    y, rest = k.split("-", 1)
    return f"{int(y) - 1}-{rest}"


def _yoy_level(d: dict) -> dict:
    """Рост г/г (%) уровня ряда {'YYYY-MM': v}."""
    return {k: round((v / d[p] - 1) * 100, 2) for k, v in d.items()
            if (p := _prev_year(k)) in d and d[p]}


def _q_rows(rows: list) -> dict:
    """Лист «строка лет → строка кварталов → строки показателей» → {метка: {'YYYY-Qn': v}}."""
    qi = next((i for i, r in enumerate(rows)
               if sum(1 for c in r if isinstance(c, str) and "квартал" in c.lower()) >= 3), None)
    if not qi:
        return {}
    yrow, qrow = rows[qi - 1], rows[qi]
    qmap = {"i": 1, "ii": 2, "iii": 3, "iv": 4}
    cols, cur = {}, None
    for j in range(1, len(qrow)):
        y = _year(yrow[j]) if j < len(yrow) else None
        if y:
            cur = y
        q = str(qrow[j]).strip().lower().split(" ")[0] if qrow[j] else ""
        if cur and q in qmap:
            cols[j] = f"{cur}-Q{qmap[q]}"
    res = {}
    for r in rows[qi + 1:]:
        lbl = re.sub(r"\s*\d\)$", "", str(r[0] or "")).strip().lower()
        if lbl and lbl not in res:
            s = {c: _f(r[j]) for j, c in cols.items() if j < len(r) and _f(r[j]) is not None}
            if s:
                res[lbl] = s
    return res


# ─── Росстат: реальный сектор ───────────────────────────────────────────────────────
def rs_cpi_components() -> dict:
    """ИПЦ по компонентам (ipc_mes): всего / продовольственные / непродовольственные / услуги.
    м/м (% к пред. месяцу) и г/г (цепочкой 12 месяцев), с 2010."""
    b = _rs_book(_RS + "/statistics/price", r"ipc_mes_\d{2}-\d{4}\.xlsx")
    out = {}
    for key, sn in (("total", "01"), ("food", "02"), ("nonfood", "03"), ("services", "04")):
        mm = _ym_grid(b.get(sn, []))
        out[key] = {"mom": _ser(_mom(mm), "2010"), "yoy": _ser(_yoy_chain(mm), "2010")}
    return out


def rs_ppi() -> dict:
    """Цены производителей (Proizvoditeli_Ind_VED): промышленность / добыча / обработка /
    электроэнергия. Листы 1.x (1998-2016) + 2.x (2013→) → м/м и г/г."""
    b = _rs_book(_RS + "/statistics/price", r"Proizvoditeli_Ind_VED_\d{2}-\d{4}\.xlsx")
    series = {}
    for key, old, new in (("industry", "1.1", "2.1"), ("mining", "1.2", "2.2"),
                          ("manuf", "1.3", "2.3"), ("energy", "1.4", "2.4")):
        series[key] = {}
        for sn in (old, new):
            series[key].update(_ym_grid(b.get(sn, [])))
    # текущий год (и прочие годы после сводных листов) — в отдельных листах «в ГГГГ г.» → «N.1 к
    # предыдущему месяцу»: строка вида деятельности (код), под ней «Российская Федерация» с месяцами
    last = max((k for s in series.values() for k in s), default="0000")
    codes = {"1323500.029.31": "industry", "B": "mining", "C": "manuf", "D": "energy"}  # 1323500… = «Промышленность»
    year_sheet, cur = {}, None
    for r in b.get("Содержание", []):
        t = " ".join(str(x) for x in r if x)
        m = re.search(r"в (\d{4}) г", t)
        if m and "Индексы цен" in t and not re.match(r"\s*\d", t):
            cur = m.group(1)
        m2 = re.match(r"\s*(\d+)\.1\s.*к предыдущему месяцу", t)
        if m2 and cur:
            year_sheet[cur] = f"{m2.group(1)}.1"
    for yy, sn in sorted(year_sheet.items()):
        if f"{yy}-12" <= last or sn not in b:
            continue
        rows = b[sn]
        hi = next((i for i, r in enumerate(rows) if sum(1 for c in r if _mon(c)) >= 3), None)
        if hi is None:
            continue
        mcols = {j: _mon(c) for j, c in enumerate(rows[hi]) if _mon(c)}
        pending = None
        for r in rows[hi + 1:]:
            code = str(r[1] or "").strip() if len(r) > 1 else ""
            vals = {f"{yy}-{m:02d}": _f(r[j]) for j, m in mcols.items() if j < len(r) and _f(r[j]) is not None}
            if code in codes:
                pending = codes[code]
                if vals:                                   # формат без региональной разбивки
                    series[pending].update(vals); pending = None
            elif pending and str(r[0] or "").strip().startswith("Российская"):
                series[pending].update(vals); pending = None
            elif code:
                pending = None
    return {key: {"mom": _ser(_mom(mm), "2010"), "yoy": _ser(_yoy_chain(mm), "2010")} for key, mm in series.items()}


_IP_CODES = [("BCDE", "total", "Промышленность"), ("B", "mining", "Добыча"), ("C", "manuf", "Обработка"),
             ("D", "energy", "Электроэнергия, газ"), ("E", "water", "Водоснабжение"),
             ("10", "food", "Пищевые продукты"), ("19", "oil_ref", "Нефтепереработка"),
             ("20", "chem", "Химия"), ("21", "pharma", "Фармацевтика"), ("24", "metal", "Металлургия"),
             ("25", "metal_prod", "Готовые металлоизделия"), ("26", "electronics", "Электроника"),
             ("27", "el_equip", "Электрооборудование"), ("28", "machinery", "Машины и оборудование"),
             ("29", "autos", "Автомобили"), ("30", "transport_eq", "Прочие трансп. средства")]


def _ip_sheet(rows: list) -> dict:
    """Лист ind_baza: строка «Наименование ОКВЭД2 | Код | 2020 год …», строка месяцев,
    строки видов деятельности (код во 2-й колонке) → {код: {'YYYY-MM': v}}."""
    hi = next(i for i, r in enumerate(rows) if r and str(r[0] or "").strip().startswith("Наименование"))
    yrow, mrow = rows[hi], rows[hi + 1]
    cols, cur = {}, None
    for j in range(2, len(mrow)):
        y = _year(yrow[j]) if j < len(yrow) else None
        if y:
            cur = y
        m = _mon(mrow[j])
        if cur and m:
            cols[j] = f"{cur}-{m:02d}"
    res = {}
    for r in rows[hi + 2:]:
        c = r[1] if len(r) > 1 else None
        code = str(int(c)) if isinstance(c, float) and c.is_integer() else str(c or "").strip()
        if code and code not in res:
            res[code] = {cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None}
    return res


def rs_industrial() -> dict:
    """Индекс промпроизводства, % г/г (ind_baza, лист «к соотв. месяцу пред. года»):
    итог, 4 раздела и ключевые отрасли обработки. Старый файл (2018=100) + новый (2023=100)."""
    pg = _RS + "/enterprise_industrial"
    merged: dict = {}
    for pat in (r"ind_baza_2018_\d{2}-\d{4}\.xlsx", r"ind_baza_2023_\d{2}-\d{4}\.xlsx"):
        try:
            b = _rs_book(pg, pat, only=["2"])
            for code, s in _ip_sheet(b["2"]).items():
                merged.setdefault(code, {}).update(s)
        except Exception as e:  # noqa: BLE001
            log.warning("rs_industrial %s: %s", pat, e)
    out = {"labels": {key: lbl for _, key, lbl in _IP_CODES}}
    for code, key, _ in _IP_CODES:
        out[key] = [{"date": k, "val": round(v - 100, 2)} for k, v in sorted(merged.get(code, {}).items())]
    return out


def rs_ivbo() -> list:
    """Индекс выпуска товаров и услуг по базовым видам деятельности, % г/г, помесячно с 2014 —
    месячный прокси ВВП (Росстат, IVBO_OKVED2)."""
    b = _rs_book(_RS + "/statistics/accounts", r"IVBO_OKVED2_\d{2}-\d{4}\.xlsx", only=["1"])
    return _ser(_mom(_ym_grid(b["1"])), "2014")


def rs_retail() -> dict:
    """Оборот розничной торговли, млрд руб/мес (Oborot_ГГГГ.xls — история + Oborot_m — текущий
    год): всего / продовольственные / непродовольственные, рост г/г номинал; плюс реальный
    (индекс физобъёма) по годам — IFO_rozn."""
    pg = _RS + "/statistics/roznichnayatorgovlya"
    tot, food, nonfood = {}, {}, {}

    def eat(rows):
        cur = None
        for r in rows:
            if not r:
                continue
            m = _mon(r[0])
            y = None if m else _year(r[0])
            if y:
                cur = y
                continue
            if cur and m and len(r) > 1 and _f(r[1]) is not None:
                k = f"{cur}-{m:02d}"
                tot[k] = round(_f(r[1]) / 1000, 1)
                if len(r) > 3 and _f(r[2]) is not None and _f(r[3]) is not None:
                    food[k], nonfood[k] = round(_f(r[2]) / 1000, 1), round(_f(r[3]) / 1000, 1)

    for pat, sn in ((r"Oborot_\d{4}\.xls", "2"), (r"Oborot_m_\d{2}-\d{4}\.xlsx", "1")):
        try:
            eat(_rs_book(pg, pat, only=[sn])[sn])
        except Exception as e:  # noqa: BLE001
            log.warning("rs_retail %s: %s", pat, e)
    real = {}
    try:
        for r in _rs_book(pg, r"IFO_rozn_\d{4}\.xls", only=["1"])["1"]:
            y = _year(r[0]) if r else None
            v = _f(r[1]) if y and len(r) > 1 else None
            if y and v and str(y) not in real:
                real[str(y)] = round(v - 100, 2)
    except Exception as e:  # noqa: BLE001
        log.warning("rs_retail real: %s", e)
    return {"total": _ser(tot, "2010"), "food": _ser(food, "2010"), "nonfood": _ser(nonfood, "2010"),
            "yoy": _ser(_yoy_level(tot), "2011"), "real_annual": _ser(real, "2000")}


def rs_wages() -> dict:
    """Среднемесячная номинальная начисленная зарплата, руб (tab1-zpl), и её рост г/г.
    Реальный рост считаем на фронте (номинал г/г против ИПЦ г/г)."""
    b = _rs_book(_RS + "/labor_market_employment_salaries", r"tab1-zpl_\d{2}-\d{4}\.xlsx")
    rows = next(iter(b.values()))
    hi = next(i for i, r in enumerate(rows) if sum(1 for c in r if _mon(c)) >= 6)
    mcols = {j: _mon(c) for j, c in enumerate(rows[hi]) if _mon(c)}
    nom = {}
    for r in rows[hi + 1:]:
        y = _year(r[0]) if r else None
        if not y:
            continue
        for j, m in mcols.items():
            v = _f(r[j]) if j < len(r) else None
            if v:
                nom[f"{y}-{m:02d}"] = v
    return {"nominal": _ser(nom, "2010"), "yoy": _ser(_yoy_level(nom), "2011")}


def rs_real_income() -> dict:
    """Реальные и реальные располагаемые денежные доходы населения, % г/г, поквартально
    (urov_12kv). Ячейки со сносками вида '98,6 2)'."""
    b = _rs_book(_RS + "/folder/13397", r"urov_12kv_[^\"/]*\.xlsx")
    rows = b.get("РРДД_РДД") or list(b.values())[-1]
    cur, ri, rdi = None, {}, {}
    for r in rows:
        if not r or r[0] is None:
            continue
        s = str(r[0]).strip().lower()
        y = _year(s)
        if y and "год" in s:
            cur = y
            continue
        m = re.match(r"(\d)\s*квартал", s)
        if cur and m:
            k = f"{cur}-Q{m.group(1)}"
            a = _f(r[1]) if len(r) > 1 else None
            c = _f(r[3]) if len(r) > 3 else None
            if a:
                ri[k] = round(a - 100, 2)
            if c:
                rdi[k] = round(c - 100, 2)
    return {"real_income": _ser(ri), "real_disposable": _ser(rdi)}


def rs_unemployment() -> dict:
    """Уровень безработицы 15+ (Trud-3_15-s): по РФ по годам (лист 2) и последние периоды
    «в среднем за три месяца» (лист 4). Помесячного ряда в открытых файлах Росстата нет."""
    b = _rs_book(_RS + "/labour_force", r"Trud-3_15-s\.xlsx", only=["2", "4"])

    def rf(rows):
        hi = next(i for i, r in enumerate(rows) if sum(1 for c in r[1:] if c not in (None, "")) >= 3)
        hdr = rows[hi]
        row = next(r for r in rows[hi + 1:] if r and str(r[0] or "").strip().startswith("Российская"))
        return [(str(hdr[j]).strip(), _f(row[j])) for j in range(1, len(hdr))
                if hdr[j] not in (None, "") and j < len(row) and _f(row[j]) is not None]

    annual = [{"date": str(_year(lb)), "val": v} for lb, v in rf(b["2"]) if _year(lb)]
    rolling = {}
    if "4" in b:  # подписи «апрель - июнь 2026» → конец окна '2026-06' (скользящее среднее 3 мес)
        try:
            for lb, v in rf(b["4"]):
                ms = [_MFULL[w] for w in re.findall(r"[а-яё]+", lb.lower()) if w in _MFULL]
                ys = re.findall(r"(?:19|20)\d{2}", lb)
                if ms and ys:
                    rolling[f"{ys[-1]}-{ms[-1]:02d}"] = v
        except StopIteration:
            pass
    out = {"annual": annual, "rolling3m": _ser(rolling)}
    try:
        out["monthly"] = _rs_unemp_monthly()
    except Exception as e:  # noqa: BLE001
        log.warning("unemployment monthly: %s", e)
    return out


_OSN_CACHE = os.path.join(os.path.dirname(__file__), "data", "osn_cache.json")


def _rs_unemp_monthly() -> list:
    """Помесячный уровень безработицы (МОТ) из доклада Росстата «Социально-экономическое положение
    России» (osn-MM-ГГГГ.pdf, табл. «Динамика численности рабочей силы»: прошлый и текущий год).
    В xlsx Росстата помесячного ряда нет — только скользящие 3 мес. PDF большой → разбор кэшируется
    по имени файла; накопленные месяцы из прошлых докладов сохраняются."""
    import io
    import json
    import pdfplumber
    t = _page(f"{_RS}/compendium/document/50801")
    links = sorted(set(re.findall(r'href="(/storage/mediabank/osn-(\d{2})-(\d{4})\.pdf)"', t)), key=lambda x: (x[2], x[1]))
    if not links:
        return []
    href = links[-1][0]
    try:
        with open(_OSN_CACHE, encoding="utf-8") as f:
            cache = json.load(f)
    except Exception:  # noqa: BLE001
        cache = {}
    acc = cache.get("months", {})
    if cache.get("src") != href:
        pdf = _rosstat_get(_RS + href)
        with pdfplumber.open(io.BytesIO(pdf)) as doc:
            n = len(doc.pages)
            order = list(range(n // 2, n)) + list(range(0, n // 2))     # раздел «Занятость» — во второй половине
            start = next((i for i in order if "ДИНАМИКА ЧИСЛЕННОСТИ РАБОЧЕЙ СИЛЫ" in (doc.pages[i].extract_text() or "")), None)
            if start is None:
                return _ser(acc)
            text = "\n".join((doc.pages[j].extract_text() or "") for j in range(start, min(start + 2, n)))
        text = text.split("Таблица 2")[0]
        year = None
        for line in text.split("\n"):
            ym = re.match(r"\s*(\d{4}) г\.", line)
            if ym:
                year = ym.group(1)
                continue
            mm = re.match(r"\s*([А-Я][а-я]+)\s+((?:[\d,]+\s+){6}[\d,]+)", line)
            if year and mm and mm.group(1).lower() in _MFULL:
                nums = mm.group(2).split()
                acc[f"{year}-{_MFULL[mm.group(1).lower()]:02d}"] = float(nums[6].replace(",", "."))
        try:
            os.makedirs(os.path.dirname(_OSN_CACHE), exist_ok=True)
            _dump_json(_OSN_CACHE, {"src": href, "months": acc})
        except Exception as e:  # noqa: BLE001
            log.warning("osn cache save: %s", e)
    return _ser(acc)


def rs_housing() -> dict:
    """Цены на жильё (dinamika_1998-ГГГГ): индекс «все типы квартир», % дек/дек — первичный
    (лист 1) и вторичный (лист 2) рынок, по годам."""
    b = _rs_book(_RS + "/statistics/price", r"dinamika_1998-\d{4}\.xlsx", only=["1", "2"])
    out = {}
    for key, sn in (("primary", "1"), ("secondary", "2")):
        s = {}
        for r in b.get(sn, []):
            y = _year(r[0]) if r else None
            v = _f(r[1]) if y and len(r) > 1 else None
            if y and v and str(y) not in s:
                s[str(y)] = round(v - 100, 2)
        out[key] = _ser(s)
    return out


def rs_confidence() -> dict:
    """Индекс предпринимательской уверенности по отраслям, помесячно с 2017 (IPU_kp, лист 1)
    и индекс потребительской уверенности, поквартально с 2006 (D-potr_sektor)."""
    pg = _RS + "/leading_indicators"
    out: dict = {"business": {}, "business_long": {}, "consumer": []}

    def sectors(rows):
        """Блоки «отрасль | год | 12 месяцев» → {отрасль: [{date,val}]}."""
        cur, acc = None, {}
        for r in rows:
            if not r or len(r) < 3:
                continue
            y = _year(r[1])
            if isinstance(r[0], str) and r[0].strip() and y:
                cur = re.sub(r"\s+", " ", re.sub(r"\s*\d\)$", "", r[0].strip()))
            if cur and y:
                for j in range(2, min(14, len(r))):
                    v = _f(r[j])
                    if v is not None:
                        acc.setdefault(cur, {})[f"{y}-{j - 1:02d}"] = v
        return {k: _ser(v) for k, v in acc.items()}

    try:
        b = _rs_book(pg, r"IPU_kp_\d{4}-\d{2}\.xlsx", only=["1", "4"])
        out["business"] = sectors(b.get("1", []))        # новая методология, с 2017
        out["business_long"] = sectors(b.get("4", []))   # старая методология, с 2006
    except Exception as e:  # noqa: BLE001
        log.warning("rs_confidence business: %s", e)
    try:
        rows = next(iter(_rs_book(pg, r"D-potr_sektor_[^\"/]*\.xls").values()))
        s = {}
        for r in rows:
            y = _year(r[1]) if len(r) > 1 else None
            if y:
                for q in range(1, 5):
                    v = _f(r[1 + q]) if len(r) > 1 + q else None
                    if v is not None:
                        s[f"{y}-Q{q}"] = v
        out["consumer"] = _ser(s)
    except Exception as e:  # noqa: BLE001
        log.warning("rs_confidence consumer: %s", e)
    return out


_USE = [("gdp", "валовой внутренний продукт"), ("consumption", "расходы на конечное потребление"),
        ("hh", "домашних хозяйств"), ("gov", "государственного управления"),
        ("accumulation", "валовое накопление"), ("gfcf", "валовое накопление основного капитала"),
        ("inventories", "изменение запасов материальных оборотных средств"),
        ("net_exports", "чистый экспорт"), ("exports", "экспорт"), ("imports", "импорт")]


def rs_gdp_use() -> dict:
    """ВВП по использованию, поквартально (GDP-quarters-of-use .xls): номинал млрд руб
    (лист 2, с 2011) и реальный рост % г/г (листы 12+13, с 1996) по элементам."""
    b = _rs_book(_RS + "/statistics/accounts", r"GDP-quarters-of-use-[^\"/]*\.xls", only=["2", "12", "13"])

    def merged(sheets):
        m: dict = {}
        for sn in sheets:
            for lbl, s in _q_rows(b.get(sn, [])).items():
                m.setdefault(lbl, {}).update(s)
        return m

    nom, real = merged(["2"]), merged(["12", "13"])
    out: dict = {"nominal": {}, "real": {}}
    for key, tgt in _USE:
        hit = next((v for k, v in nom.items() if k.startswith(tgt) and (tgt != "валовое накопление" or k == tgt)), None)
        if hit:
            out["nominal"][key] = _ser(hit)
        hit = next((v for k, v in real.items() if k.startswith(tgt) and (tgt != "валовое накопление" or k == tgt)), None)
        if hit:
            out["real"][key] = _ser({k: round(v - 100, 2) for k, v in hit.items()})
    return out


# ─── ЦБ: деньги, ликвидность, внешний сектор ───────────────────────────────────────
def _dmy2iso(s: str) -> str:
    return f"{s[6:10]}-{s[3:5]}-{s[:2]}"


def cbr_reserves() -> dict:
    """Международные резервы РФ, млрд $: помесячно с 2008 (всего / валютные / монетарное
    золото, hd_base/mrrf_m) и понедельно за 3 года (mrrf_7d)."""
    now = datetime.now()
    rows = _hd_rows(f"{_CBR}/hd_base/mrrf/mrrf_m/?UniDbQuery.Posted=True"
                    f"&UniDbQuery.From=01.2008&UniDbQuery.To={now:%m.%Y}")
    tot, fxr, gold = {}, {}, {}
    for r in rows:  # [дата, всего, валютные, ин.валюта, СДР, МВФ, золото] млн $
        k = f"{r[0][6:10]}-{r[0][3:5]}"
        for d, j in ((tot, 1), (fxr, 2), (gold, 6)):
            v = _num(r[j]) if j < len(r) else None
            if v is not None:
                d[k] = round(v / 1000, 1)
    weekly = []
    try:
        wr = _hd_rows(f"{_CBR}/hd_base/mrrf/mrrf_7d/?UniDbQuery.Posted=True"
                      f"&UniDbQuery.From={now - timedelta(days=1100):%d.%m.%Y}&UniDbQuery.To={now:%d.%m.%Y}")
        weekly = sorted(({"date": _dmy2iso(r[0]), "val": _num(r[1])} for r in wr if len(r) > 1 and _num(r[1]) is not None),
                        key=lambda x: x["date"])
    except Exception as e:  # noqa: BLE001
        log.warning("cbr_reserves weekly: %s", e)
    return {"total": _eom(_ser(tot)), "fx": _eom(_ser(fxr)), "gold": _eom(_ser(gold)), "weekly": weekly}   # помесячные «на 1-е число» → конец пред. месяца


def cbr_keyrate_hist() -> list:
    """Ключевая ставка — точки изменения с 17.09.2013 [{date:'YYYY-MM-DD', val}] + сегодня."""
    now = datetime.now()
    rows = _hd_rows(f"{_CBR}/hd_base/KeyRate/?UniDbQuery.Posted=True"
                    f"&UniDbQuery.From=17.09.2013&UniDbQuery.To={now:%d.%m.%Y}")
    pts = sorted((_dmy2iso(r[0]), _num(r[1])) for r in rows if len(r) > 1 and _num(r[1]) is not None)
    out, last = [], None
    for d, v in pts:
        if v != last:
            out.append({"date": d, "val": v})
            last = v
    if pts and out[-1]["date"] != pts[-1][0]:
        out.append({"date": pts[-1][0], "val": pts[-1][1]})
    return out


def cbr_liquidity() -> list:
    """Структурный дефицит(+)/профицит(−) ликвидности банковского сектора, млрд руб,
    среднее за месяц (hd_base/bliquidity, дневные данные с 2016)."""
    now = datetime.now()
    rows = _hd_rows(f"{_CBR}/hd_base/bliquidity/?UniDbQuery.Posted=True"
                    f"&UniDbQuery.From=01.01.2016&UniDbQuery.To={now:%d.%m.%Y}")
    acc: dict = {}
    for r in rows:
        v = _num(r[1]) if len(r) > 1 else None
        if v is not None:
            acc.setdefault(f"{r[0][6:10]}-{r[0][3:5]}", []).append(v)
    return [{"date": k, "val": round(sum(v) / len(v), 1)} for k, v in sorted(acc.items())]


def cbr_money_base() -> dict:
    """Денежная база в широком определении, млрд руб, помесячно (mb_bd.xlsx): всего,
    наличные, корсчета, обязательные резервы, депозиты КО в ЦБ, облигации ЦБ у КО."""
    rows = next(iter(_cbr_book(f"{_CBR}/vfs/statistics/ms/mb_bd.xlsx").values()))
    keys = ("total", "cash", "corr", "reserves", "deposits", "bonds")
    out: dict = {k: {} for k in keys}
    for r in rows:
        d = r[0] if r else None
        if isinstance(d, datetime):
            k = d.strftime("%Y-%m")
        elif isinstance(d, str) and re.match(r"\d{2}\.\d{2}\.\d{4}", d.strip()):
            k = f"{d.strip()[6:10]}-{d.strip()[3:5]}"
        else:
            continue
        for i, key in enumerate(keys, 1):
            v = _f(r[i]) if i < len(r) else None
            if v is not None:
                out[key][k] = v
    return {k: _eom(_ser(v, "2010")) for k, v in out.items()}      # «на 1-е число» → конец пред. месяца


_CREDIT_ROWS = {"credit": "кредит экономике", "corp": "требования к организациям",
                "corp_loans": "корпоративные кредиты", "corp_bonds": "корпоративные облигации",
                "households": "требования к населению*", "retail_loans": "кредиты физическим лицам",
                "m2x": "широкая денежная масса м2х", "m2": "денежная масса м2", "m0": "наличные деньги м0",
                "dep_hh": "депозиты населения", "dep_corp": "депозиты организаций"}


def cbr_credit_economy() -> dict:
    """«Денежные агрегаты и кредит экономике» (credit_m2x.xlsx, лист YOY): годовые темпы
    прироста, % — кредит экономике, корпкредиты, корпоблигации, кредиты населению, М2Х…"""
    url = _link(_CBR + "/statistics/macro_itm/dkfs/", r"credit_m2x\.xlsx", rosstat=False)
    rows = _cbr_book(url, only=["YOY"])["YOY"]
    cols = {j: c.strftime("%Y-%m") for j, c in enumerate(rows[0]) if isinstance(c, datetime)}
    out = {}
    for r in rows[1:]:
        lbl = str(r[0] or "").strip().lower()
        for key, t in _CREDIT_ROWS.items():
            if lbl == t and key not in out:
                out[key] = _ser({cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None})
    return {k: _eom(v) for k, v in out.items()}      # шапка «на 1-е число» → конец пред. месяца


_HH_ROWS = [("cash", "наличная валюта"), ("deposits", "депозиты"), ("bonds", "долговые ценные бумаги"),
            ("equity", "акции и паи и акции инвестиционных фондов"),
            ("insurance", "страховые и пенсионные резервы"),
            ("loans", "кредиты кредитных организаций - всего"), ("mfo", "займы, полученные от микрофинансовых")]


def cbr_households() -> dict:
    """Финансовые активы и обязательства домохозяйств, млрд руб, помесячно с 2018
    (households_bm.xlsx): наличные, депозиты, облигации, акции+паи, страх./пенс. резервы;
    обязательства — кредиты банков, займы МФО."""
    rows = _cbr_book(f"{_CBR}/vfs/statistics/households/households_bm.xlsx", only=["Балансы"])["Балансы"]
    hi = next(i for i, r in enumerate(rows) if sum(1 for c in r if isinstance(c, datetime)) >= 6)
    cols = {j: c.strftime("%Y-%m") for j, c in enumerate(rows[hi]) if isinstance(c, datetime)}
    out = {}
    for r in rows[hi + 1:]:
        lbl = re.sub(r"\(\d+\)|\*", "", str(r[0] or "")).strip().lower()
        for key, t in _HH_ROWS:
            if key not in out and (lbl == t or (len(t) > 20 and lbl.startswith(t))):
                out[key] = _ser({cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None})
    return {k: _eom(v) for k, v in out.items()}      # балансы «на 1-е число» → конец пред. месяца


_OKVED_SHORT = {"A": "Сельское хозяйство", "B": "Добыча", "C": "Обработка", "D": "Электроэнергетика",
                "F": "Строительство", "G": "Торговля", "H": "Транспорт", "I": "Гостиницы, общепит",
                "J": "ИТ и связь", "L": "Недвижимость", "M": "Проф. и научная", "N": "Административная"}


def _okved_short(lbl: str) -> str:
    m = re.match(r"^([A-Z])\s*-\s*", lbl)
    if m:
        return _OKVED_SHORT.get(m.group(1), lbl[:24])
    low = lbl.lower()
    if low.startswith(("нефинансовые", "итого")):
        return "Все организации"
    if low.startswith("прочие"):
        return "Прочие"
    return lbl[:30]


def cbr_nfs() -> dict:
    """Финансовое состояние нефинансовых организаций по отраслям, по годам (ЦБ,
    Stat_fin_cond_nfs.xlsx): лист 10 — показатели финустойчивости (стоимость долга,
    финансовый рычаг…), лист 5 — обязательства, трлн руб."""
    b = _cbr_book(f"{_CBR}/vfs/statistics/macro_itm/Stat_fin_cond_nfs.xlsx", only=["5", "10"])

    def blocks(rows):
        res, years = [], None
        for r in rows:
            cells = list(r[1:])
            lbl = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
            if sum(1 for c in cells if isinstance(c, datetime)) >= 3:
                # колонка «01.01.2025» = годовая отчётность за 2024 → подписываем отчётным годом
                years = {j + 1: c.year - 1 for j, c in enumerate(cells) if isinstance(c, datetime)}
                continue
            has_num = any(_f(c) is not None for c in cells[1:])
            if lbl and not has_num:
                if not lbl.startswith("*"):
                    res.append({"title": re.sub(r"\s+", " ", lbl), "rows": {}})
                    years = None
                continue
            if years and res and lbl:
                s = {str(y): _f(r[j]) for j, y in years.items() if j < len(r) and _f(r[j]) is not None}
                if s:
                    res[-1]["rows"][_okved_short(lbl)] = _ser(s)
        return [x for x in res if x["rows"]]

    return {"stability": blocks(b.get("10", [])), "liabilities": blocks(b.get("5", []))}


def cbr_ext_debt() -> dict:
    """Внешний долг РФ, млрд $, поквартально с 2003 (debt_new.xlsx): всего, госуправление,
    ЦБ и банки, прочие секторы. Дата — на начало квартала (= конец предыдущего)."""
    url = _link(_CBR + "/statistics/macro_itm/external_sector/ed/", r"debt_new\.xlsx", rosstat=False)
    rows = next(iter(_cbr_book(url).values()))
    hi = next(i for i, r in enumerate(rows) if sum(1 for c in r if isinstance(c, datetime)) >= 6)
    cols = {j: c.strftime("%Y-%m") for j, c in enumerate(rows[hi]) if isinstance(c, datetime)}
    want = {"total": "всего", "gov": "органы государственного управления",
            "banks": "центральный банк и банки", "other": "прочие секторы"}
    out = {}
    for r in rows[hi + 1:]:
        lbl = str(r[0] or "").strip().lower()
        for key, t in want.items():
            if lbl == t and key not in out:
                out[key] = _ser({cols[j]: round(_f(r[j]) / 1000, 1) for j in cols
                                 if j < len(r) and _f(r[j]) is not None})
    # предварительная оценка ЦБ на следующую квартальную дату (debt_est.xlsx) выходит на ~1,5 мес
    # раньше полного файла → дописываем её с флагом est
    try:
        rows = next(iter(_cbr_book(_CBR + "/vfs/statistics/credit_statistics/debt/debt_est.xlsx").values()))
        hi = next(i for i, r in enumerate(rows) if sum(1 for c in r if isinstance(c, datetime)) >= 3)
        ecols = {}
        for j, c in enumerate(rows[hi]):
            if isinstance(c, datetime):
                ecols[j] = (c.strftime("%Y-%m"), False)
            elif isinstance(c, str) and re.match(r"\s*\d{2}\.\d{2}\.\d{4}", c):
                d0 = re.match(r"\s*(\d{2})\.(\d{2})\.(\d{4})", c)
                ecols[j] = (f"{d0.group(3)}-{d0.group(2)}", "оценк" in c.lower())
        upd = next((re.search(r"(\d{1,2} [а-я]+ \d{4})", str(r[0])).group(1) for r in rows if r and "последнего обновления" in str(r[0]) and re.search(r"\d{1,2} [а-я]+ \d{4}", str(r[0]))), None)
        for r in rows[hi + 1:]:
            lbl = str(r[0] or "").strip().lower()
            for key, t in want.items():
                if lbl == t and key in out:
                    have = {x["date"] for x in out[key]}
                    for j, (dt, is_est) in sorted(ecols.items(), key=lambda kv: kv[1][0]):
                        v = _f(r[j]) if j < len(r) else None
                        if v is not None and dt not in have and dt > max(have):
                            out[key].append({"date": dt, "val": round(v / 1000, 1), **({"est": True} if is_est else {})})
        if upd:
            out["est_updated"] = upd
    except Exception as e:  # noqa: BLE001
        log.warning("ext_debt est: %s", e)
    return {k: (_eom(v) if isinstance(v, list) else v) for k, v in out.items()}   # «на 1 июля» → конец II кв.


def cbr_trade() -> dict:
    """Внешняя торговля товарами по методологии платёжного баланса, млрд $/мес, с 1997
    (trade.xls, лист «Ежемесячные»): экспорт, импорт, сальдо."""
    rows = _cbr_book(f"{_CBR}/vfs/statistics/credit_statistics/trade/trade.xls", only=["Ежемесячные"])["Ежемесячные"]
    cur, ex, im, bal = None, {}, {}, {}
    for r in rows:
        if not r:
            continue
        y = _year(r[0]) if r[0] not in (None, "") else None
        if y:
            cur = y
        m = _mon(r[1]) if len(r) > 1 else None
        if cur and m:
            k = f"{cur}-{m:02d}"
            for d, j in ((ex, 2), (im, 8), (bal, 14)):
                v = _f(r[j]) if j < len(r) else None
                if v is not None:
                    d[k] = round(v / 1000, 2)
    return {"exports": _ser(ex), "imports": _ser(im), "balance": _ser(bal)}


def cbr_rate_terms(y1: int = 2014) -> dict:
    """Ставки по срокам (data-service, в рублях): вклады физлиц, кредиты юрлицам и физлицам
    (все сроки), плюс ипотека / автокредиты / МСП. {группа: {название срока: [{date,val}]}}."""
    out = {}
    for key, pub, ds in (("dep_retail", 18, 37), ("dep_corp", 18, 38), ("cred_corp", 14, 25),
                         ("cred_retail", 14, 27), ("cred_msp", 14, 26), ("cred_auto", 14, 28),
                         ("cred_mortgage", 14, 29)):
        try:
            names, cols = _ds_multi(pub, ds, 2, y1)
            if not cols:   # ипотека: разрез «В рублях / В ин. валюте» — без measureId
                names, cols = _ds_multi(pub, ds, None, y1)
            out[key] = {re.sub(r"\s+", " ", names[c]).strip(): s for c, s in cols.items() if c in names and s}
        except Exception as e:  # noqa: BLE001
            log.warning("cbr_rate_terms %s: %s", key, e)
            out[key] = {}
    return out


def cbr_issuance(y1: int = 2015) -> dict:
    """Объём выданных кредитов за месяц, млрд руб (data-service, РФ): физлицам, юрлицам."""
    out = {}
    for key, pub, ds in (("retail", 20, 41), ("corp", 22, 49)):
        try:
            n, c = _ds_multi(pub, ds, 22, y1)
            s = _pick(n, c, "всего") or next(iter(c.values()), [])
            out[key] = [{"date": x["date"], "val": round(x["val"] / 1000, 1)} for x in s]
        except Exception as e:  # noqa: BLE001
            log.warning("cbr_issuance %s: %s", key, e)
            out[key] = []
    return out


def minfin_gov_debt() -> dict:
    """Государственный внутренний долг РФ, млрд руб, на 1-е число месяца, с 2009 (Минфин,
    «Таблицы ежемесячных значений объёма госвнутреннего долга», .xls — лист на год): всего,
    госгарантии и ценные бумаги/кредиты (= всего − гарантии, в основном ОФЗ).
    Замена ряду «обслуживание госдолга»: его Минфин помесячно не раскрывает с 2022 года."""
    import xlrd
    page = "https://minfin.gov.ru/ru/document?id_4=93479-tablitsy_ezhemesyachnykh_znachenii_obema_gosudarstvennogo_vnutrennego_dolga_rossiiskoi_federatsii"
    import requests
    hdr = {"User-Agent": _UA, "Referer": "https://minfin.gov.ru/"}
    html = requests.get(page, headers=hdr, timeout=40).text
    m = re.search(r'href="([^"]*Obem_gosdolga[^"]*\.xls)"', html)
    if not m:
        raise RuntimeError("нет файла объёма госдолга на странице Минфина")
    url = m.group(1) if m.group(1).startswith("http") else "https://minfin.gov.ru" + m.group(1)
    wb = xlrd.open_workbook(file_contents=requests.get(url, headers=hdr, timeout=60).content)
    tot, gar = {}, {}
    for sh in wb.sheets():
        rows = [sh.row_values(i) for i in range(sh.nrows)]
        cols = {}
        for r in rows:
            # даты ищем ТОЛЬКО в шапке — до первой строки данных («…долг…», «…гарант…»): иначе в
            # январском листе (мало колонок) сумма долга 36 000–60 000 млрд читалась бы как дата Excel
            if any(k in str(r[0] or "").lower() for k in ("долг", "гарант")) and \
                    any(isinstance(c, float) and c > 0 for c in r[1:]) and cols:
                break
            for j, c in enumerate(r):
                if isinstance(c, float) and 36000 < c < 60000:          # дата Excel
                    dt = xlrd.xldate_as_datetime(c, wb.datemode)
                    cols[j] = dt.strftime("%Y-%m")
                elif isinstance(c, str) and re.match(r"\s*\d{2}\.\d{2}\.\d{4}", c):
                    s = c.strip()
                    cols[j] = f"{s[6:10]}-{s[3:5]}"
            if len(cols) >= 12:
                break
        got_total = False
        for r in rows:
            lbl = str(r[0] or "").lower()
            vals = {cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None}
            if not vals:
                continue
            if "гарант" in lbl:
                gar.update(vals)
            elif not got_total and ("долг" in lbl):
                tot.update(vals)
                got_total = True
    sec = {k: round(v - gar.get(k, 0), 1) for k, v in tot.items()}
    return {"total": _eom(_ser(tot)), "guarantees": _eom(_ser(gar)), "securities": _eom(_ser(sec))}   # «на 1-е число» → конец пред. месяца


# ─── Базовая / устойчивая инфляция: PDF «Инфляция в России» + Excel ЦБ ─────────────
_CPD_CACHE = os.path.join(os.path.dirname(__file__), "data", "cpd_cache.json")
_CPD_VER = 6         # поднять при изменении парсера → кэш выпусков пересоберётся
_CPD_LABELS = [       # порядок важен: более длинные подписи раньше
    ("медиана (с учетом весов)", "median_w"), ("медиана (без учета весов)", "median_u"),
    ("медиана¹", "median_w"), ("медиана²", "median_u"), ("медиана", "median_w"),
    ("ипц, все товары и услуги", "cpi"), ("все товары и услуги", "cpi"), ("инфляция", "cpi"),
    ("без плодоовощной продукции, нефтепродуктов и жку", "cpi_ex"), ("без плодоовощей, нефтепродуктов и жку", "cpi_ex"),
    ("базовый ипц без туризма", "core_ex_tour"), ("базовый ипц", "core"), ("базовая инфляция", "core"),
    ("прирост цен без 20%", "ex20"),
    ("продовольственные товары без плодоовощ", "food_ex"), ("продовольственные товары", "food"),
    ("плодоовощная продукция", "fruitveg"),
    ("непродовольственные товары без нефтепродуктов", "nonfood_ex"), ("непродовольственные товары", "nonfood"),
    ("услуги без жку", "services_ex"), ("услуги", "services"), ("среднегодовая инфляция", "avg_annual"),
]
_CPD_NUM = re.compile(r"(?<![\d,])-?\d+,\d+")


def _cpd_parse(pdf: bytes, y: int, m: int) -> dict:
    """Таблица «Динамика потребительских цен» из PDF выпуска за месяц (y, m) →
    {'yoy'|'saar': {ключ: {'YYYY-MM': v}}}. Три поколения формата (см. docstring cbr_core_inflation)."""
    import pdfplumber
    out: dict = {"yoy": {}, "saar": {}}
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        for pg in doc.pages[:6]:
            t = (pg.extract_text() or "").replace("‑", "-").replace("−", "-").replace("–", "-")
            low_all = t.lower()
            flat = re.sub(r"\s+", "", low_all)
            if "базов" not in low_all or not ("г/г" in flat or "годовыеприросты" in flat
                                              or "ксоответствующемупериоду" in flat):
                continue
            lines = t.split("\n")
            old = "годовые приросты" in low_all        # формат 2020–11.2021: «июнь(пр.г) январь … июнь»
            has_avg = "мендерс" in low_all             # «В среднем за 3 мес» (перевёрнутый текст) — лишняя колонка
            if old:
                # колонки: «июнь(год назад) январь … июнь» ИЛИ просто 6 месяцев подряд (бывает через
                # Новый год) → годы расставляем с конца: при «откате» месяца назад — минус год
                hdr = next((l for l in lines if sum(1 for w in l.lower().split() if w in _MFULL) >= 3), "")
                ms = [_MFULL[w] for w in hdr.lower().split() if w in _MFULL]
                months, yy = [], (y if ms and ms[-1] <= m else y - 1)
                for i in range(len(ms) - 1, -1, -1):
                    if i < len(ms) - 1 and ms[i] >= ms[i + 1]:
                        yy -= 1
                    months.insert(0, (yy, ms[i]))
            else:
                months = [((y * 12 + m - 1 - k) // 12, (y * 12 + m - 1 - k) % 12 + 1) for k in (2, 1, 0)]
            blk = None
            for line in lines:
                low = line.lower().strip()
                comp = re.sub(r"\s+", "", low)
                if comp.startswith(("темпприростацен,г/г", "годовыеприросты", "темпприростаценксоответствующемупериоду")):
                    blk = "yoy"; continue
                if comp.startswith("месячныеприросты,sa,среднее"):
                    blk = None; continue
                # «SAAR» / «с.к.г.» — уже в годовом выражении; «SA» — месячный прирост (пересчитываем в год).
                # Проверять SAAR раньше SA: «sa» содержится в «saar» (выпуски 12.2023–01.2024 давали 246% вместо 10%)
                if comp.startswith("темпприростаценкпредыдущемупериоду") and ("saar" in comp or "с.к.г" in comp):
                    blk = "saar"; continue
                if comp.startswith("месячныеприросты,sa") or (comp.startswith("темпприростаценкпредыдущемупериоду") and "sa" in comp):
                    blk = "sa"; continue
                if comp.startswith("темпприростаценкпредыдущемупериоду"):
                    blk = "mom"; continue
                if blk not in ("yoy", "sa", "saar"):
                    continue
                nums = _CPD_NUM.findall(line)
                if len(nums) < 3:
                    continue
                lbl = line[:line.find(nums[0])].lower()
                lbl = re.sub(r"^[\s―\-–]*(из них:)?\s*", "", lbl).replace("*", "").strip()
                key = next((k for s, k in _CPD_LABELS if lbl.startswith(s)), None)
                if not key:
                    continue
                vals = [float(n.replace(",", ".")) for n in nums]
                if old:
                    tail = vals[-len(months):] if len(vals) >= len(months) else []
                else:
                    tail = vals[-4:-1] if (has_avg and blk != "yoy") else vals[-3:]
                dst = "saar" if blk in ("sa", "saar") else "yoy"
                d = out[dst].setdefault(key, {})
                for (yy, mm), v in zip(months, tail):
                    if blk == "sa":                     # м/м SA → в годовом выражении (как «с.к.г.» ЦБ)
                        v = round(((1 + v / 100) ** 12 - 1) * 100, 1)
                    d.setdefault(f"{yy}-{mm:02d}", v)
            if out["yoy"]:
                break
    return out


def _cpd_issues() -> dict:
    """{(год, месяц): URL PDF} всех выпусков «Инфляция в России» / «Динамика потребительских цен»."""
    t = _get(f"{_CBR}/analytics/dkp/dinamic/").text
    res = {}
    for iid, yy, mm in set(re.findall(r"/analytics/dkp/dinamic/(CPD_(\d{4})-(\d{1,2}))/", t)):
        res[(int(yy), int(mm))] = ("page", iid)
    for href, yy, mm in set(re.findall(r'href="([^"]*CPD_(\d{4})-(\d{1,2})\.pdf)"', t)):
        res.setdefault((int(yy), int(mm)), ("pdf", href))
    return res


def _cpd_fetch(y: int, m: int, kind: str, ref: str) -> dict:
    if kind == "page":
        p = _get(f"{_CBR}/analytics/dkp/dinamic/{ref}/").text
        pat = re.compile(rf"CPD_{y}-0?{m}\.pdf$")
        href = next(h for h in re.findall(r'href="([^"]+\.pdf)"', p) if pat.search(h))
    else:
        href = ref
    return _cpd_parse(_get(href if href.startswith("http") else _CBR + href).content, y, m)


def cbr_core_inflation() -> dict:
    """Базовая и устойчивая инфляция.
    1) PDF ЦБ «Инфляция в России» (помесячно, с 2020): официальная базовая инфляция г/г (Росстат),
       базовый ИПЦ без туризма, медианы (с весами / без), прирост без 20% волатильных компонентов,
       всё — и г/г, и SAAR. В каждом выпуске 3 последних месяца (форматы 2022→ ) или январь…месяц
       (2020–2021) → склеиваем выпуски; значение месяца = медиана по выпускам (защита от сбоя
       текстового слоя PDF). Разобранные выпуски кэшируются на диске (PDF скачиваются один раз).
    2) Excel ЦБ «Аналитические показатели ценовой динамики» (statistics/ddkp/aipd): трендовая
       инфляция (окно 5 и 3 года), модифицированная базовая (усечение / исключение волатильных),
       сезонно сглаженный рост цен по ~50 компонентам ИПЦ (м/м SA → SAAR), с 2002."""
    import json
    import statistics
    from concurrent.futures import ThreadPoolExecutor

    out: dict = {"yoy": {}, "saar": {}, "trend": {}, "modcore": {}, "sa": {}, "yoy_chain": {}}
    # --- PDF-выпуски (кэш на диске) ---
    try:
        try:
            with open(_CPD_CACHE, encoding="utf-8") as f:
                cache = json.load(f)
            if cache.get("ver") != _CPD_VER:
                cache = {}
        except Exception:  # noqa: BLE001
            cache = {}
        issues = cache.get("issues", {})
        # ЦБ в оглавлении ссылается на страницы, которых нет (CPD_2025-4 и -5 отдают 404).
        # Такие выпуски помним неделю, иначе каждое обновление снова ходит за ними и пишет
        # предупреждение в лог. Месяцы при этом не теряются: в каждом выпуске три месяца
        import time as _time
        dead = {k: v for k, v in (cache.get("dead") or {}).items() if _time.time() - v < 7 * 86400}
        todo = [(k, v) for k, v in sorted(_cpd_issues().items())
                if f"{k[0]}-{k[1]:02d}" not in issues and f"{k[0]}-{k[1]:02d}" not in dead]

        def run(item):
            (y, m), (kind, ref) = item
            try:
                return f"{y}-{m:02d}", _cpd_fetch(y, m, kind, ref)
            except Exception as e:  # noqa: BLE001
                log.warning("cpd %s-%s: %s", y, m, e)
                return f"{y}-{m:02d}", None

        with ThreadPoolExecutor(max_workers=6) as ex:
            for iid, parsed in ex.map(run, todo):
                if parsed and parsed["yoy"]:
                    issues[iid] = parsed
                else:
                    dead[iid] = _time.time()
        try:
            os.makedirs(os.path.dirname(_CPD_CACHE), exist_ok=True)
            _dump_json(_CPD_CACHE, {"ver": _CPD_VER, "issues": issues, "dead": dead})
        except Exception as e:  # noqa: BLE001
            log.warning("cpd cache save: %s", e)
        for blk in ("yoy", "saar"):
            acc: dict = {}
            for iid in sorted(issues):
                for key, s in issues[iid].get(blk, {}).items():
                    for dt, v in s.items():
                        acc.setdefault(key, {}).setdefault(dt, []).append(v)
            if blk == "saar":
                # SAAR: сезонная модель пересчитывается в каждом выпуске → верен ПОСЛЕДНИЙ выпуск;
                # медиана — только страховка от сбоя разбора (последнее значение дальше 5 п.п.)
                pick = lambda vs: (vs[-1] if len(vs) < 3 or abs(vs[-1] - statistics.median(vs)) <= 5  # noqa: E731
                                   else statistics.median(vs))
            else:
                pick = lambda vs: statistics.median(vs) if len(vs) >= 3 else vs[-1]  # noqa: E731
            out[blk] = {key: _ser({dt: pick(vs) for dt, vs in d.items()}) for key, d in acc.items()}
        out["issues"] = len(issues)
        out["last_issue"] = max(issues) if issues else None
    except Exception as e:  # noqa: BLE001
        log.warning("cbr_core_inflation pdf: %s", e)
    # --- Excel: трендовая + модифицированная базовая, SA по компонентам ---
    aipd = _CBR + "/statistics/ddkp/aipd/"
    saar = lambda v: round(((1 + v / 100) ** 12 - 1) * 100, 2)  # noqa: E731
    try:
        rows = next(iter(_cbr_book(_link(aipd, r"est_infl\.xlsx", rosstat=False)).values()))
        hi = next(i for i, r in enumerate(rows) if sum(1 for c in r if isinstance(c, datetime)) >= 6)
        cols = {j: c.strftime("%Y-%m") for j, c in enumerate(rows[hi]) if isinstance(c, datetime)}
        for r in rows[hi + 1:]:
            lbl = str(r[0] or "").lower()
            s = {cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None}
            if not s:
                continue
            if "трендовой" in lbl:
                out["trend"]["t5" if "5-лет" in lbl else "t3"] = _ser(s)
            elif "усечения" in lbl:
                out["modcore"]["trunc"] = _ser({k: saar(v) for k, v in s.items()})
            elif "волатильн" in lbl:
                out["modcore"]["excl"] = _ser({k: saar(v) for k, v in s.items()})
    except Exception as e:  # noqa: BLE001
        log.warning("cbr_core_inflation est_infl: %s", e)
    try:
        rows = _cbr_book(_link(aipd, r"indicators_cpd\.xlsx", rosstat=False)).get("Лист1", [])
        cols = {j: c.strftime("%Y-%m") for j, c in enumerate(rows[0]) if isinstance(c, datetime)}
        for r in rows[1:]:
            # название — первая текстовая ячейка: ЦБ убирал/добавлял служебный столбец перед
            # названием (17.09.2026 названия переехали в 1-й столбец — брались числа как подписи)
            name = next((c.strip() for c in r[:3] if isinstance(c, str) and c.strip()), "")
            if not name or name.startswith(("1 -", "2 -", "3 -")):
                continue
            name = {"- без плодоовощной продукции, нефтепродуктов и ЖКХ": "ИПЦ без плодоовощей, топлива и ЖКХ",
                    "-без плодоовощной продукции": "Продовольствие без плодоовощей",
                    "- без нефтепродуктов": "Непрод. без нефтепродуктов", "- без ЖКХ": "Услуги без ЖКХ"}.get(name, name)
            raw = {cols[j]: _f(r[j]) for j in cols if j < len(r) and _f(r[j]) is not None}
            if raw:
                out["sa"][name] = _ser({k: saar(v) for k, v in raw.items()})
            # г/г по цепочке 12 SA-приростов — длинная история (с 2003) там, где нет PDF-таблиц
            ck = {"Базовый ИПЦ": "core", "Все товары и услуги": "cpi"}.get(name)
            if ck and raw:
                out["yoy_chain"][ck] = _ser(_yoy_chain({k: 100 + v for k, v in raw.items()}))
    except Exception as e:  # noqa: BLE001
        log.warning("cbr_core_inflation indicators_cpd: %s", e)
    # длинный ряд базовой инфляции г/г: расчёт по цепочке SA до первой PDF-точки, дальше официальная
    # (сверка на 77 общих месяцах: среднее расхождение 0,06 п.п., максимум 0,21)
    off, ch = out["yoy"].get("core", []), out["yoy_chain"].get("core", [])
    first = off[0]["date"] if off else "9999"
    out["core_long"] = [x for x in ch if x["date"] < first] + off
    out["core_long_from"] = first if off else None
    return out


def rosstat_weekly_items(top: int = 40, weeks: int = 26) -> dict:
    """Недельные изменения цен по товарам (Росстат, Nedel_ipc.xlsx), % за неделю, за последние
    `weeks` недель — `top` самых весомых в корзине ИПЦ (веса — Vesa-tov-KIPC). Для тепловой карты
    «по неделям»."""
    from .macro import _RU_MON
    try:
        _wurl = _link(f"{_RS}/statistics/price", r"Vesa-tov-KIPC_\d{4}-\d{4}\.xlsx")
    except Exception:  # noqa: BLE001
        _wurl = f"{_RS}/storage/mediabank/Vesa-tov-KIPC_2012-2026.xlsx"
    wv = _book(_rosstat_get(_wurl))
    wsv = wv[sorted(k for k in wv if k.isdigit())[-1]]      # веса текущего года
    weights = {str(r[1]).strip().lower(): r[2] for r in wsv if len(r) >= 3 and isinstance(r[1], str) and isinstance(r[2], (int, float))}
    from .macro import _nedel_ipc_bytes
    wb = _book(_nedel_ipc_bytes())
    series: dict = {}
    for sn in sorted(k for k in wb if k.isdigit()):                      # все годы (с 2022) — для месяцев и лет
        rows = wb[sn]
        hi = next((i for i, r in enumerate(rows)
                   if sum(1 for c in r if isinstance(c, str) and re.search(r"на\s+\d{1,2}\s+[а-яё]+", c.lower())) >= 3), None)
        if hi is None:
            continue
        cols = {}
        for j, lbl in enumerate(rows[hi]):
            mm = re.search(r"на\s+(\d{1,2})\s+([а-яё]+)", str(lbl or "").lower())
            if mm and mm.group(2) in _RU_MON:
                cols[j] = f"{sn}-{_RU_MON[mm.group(2)]:02d}-{int(mm.group(1)):02d}"
        for r in rows[hi + 1:]:
            nm = r[0].strip() if r and isinstance(r[0], str) else None
            if not nm:
                continue
            for j, dt in cols.items():
                v = r[j] if j < len(r) and isinstance(r[j], (int, float)) else None
                if v is not None:
                    series.setdefault(nm, {})[dt] = round(v - 100, 2)
    all_dates = sorted({dt for s in series.values() for dt in s})[-weeks:]

    def agg(s, n):   # недельные % → за месяц (n=7) / год (n=4): произведение (1 + x) по неделям периода
        acc: dict = {}
        for dt, v in sorted(s.items()):
            acc[dt[:n]] = acc.get(dt[:n], 1.0) * (1 + v / 100)
        return [{"date": k, "val": round((x - 1) * 100, 2)} for k, x in sorted(acc.items())]
    items = [{"name": re.sub(r",\s*[^,]+$", "", nm), "unit": nm.rsplit(",", 1)[-1].strip() if "," in nm else "",
              "weight": weights.get(nm.lower(), 0), "series": [{"date": dt, "val": s[dt]} for dt in all_dates if dt in s],
              "monthly": agg(s, 7)[-24:], "yearly": agg(s, 4)}
             for nm, s in series.items()]
    items = [x for x in items if x["series"]]
    items.sort(key=lambda x: -x["weight"])
    return {"weeks": all_dates, "items": items[:top], "total": len(items)}


# ─── Прогнозы: среднесрочный прогноз ЦБ (автоматически) и план бюджета (закон) ───────
_FC_ROWS = [  # (ключ, regex подписи строки, таблица: 1 — основные параметры, 2 — платёжный баланс)
    ("infl_dec", r"Инфляция, в %, декабрь к декабрю", 1), ("infl_avg", r"Инфляция, в среднем за год", 1),
    ("key_rate", r"Ключевая ставка, в среднем за год", 1), ("gdp", r"Валовой внутренний продукт", 1),
    ("consumption", r"Расходы на конечное потребление", 1), ("hh", r"– домашних хозяйств", 1),
    ("accumulation", r"Валовое накопление", 1), ("gfcf", r"– основного капитала", 1),
    ("exports_real", r"Экспорт", 1), ("imports_real", r"Импорт", 1),
    ("m2", r"Денежная масса в национальном определении", 1),
    ("credit", r"Требования банковской системы к экономике", 1), ("credit_corp", r"– к организациям", 1),
    ("credit_hh", r"– к населению", 1), ("mortgage", r"ипотечные жилищные кредиты", 1),
    ("current_account", r"Счет текущих операций", 2), ("trade_balance", r"Торговый баланс", 2),
    ("exports_usd", r"Экспорт", 2), ("imports_usd", r"Импорт", 2),
]
_MGEN = {m: i for i, m in enumerate(["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
                                     "сентября", "октября", "ноября", "декабря"], 1)}
_FC_TOK = re.compile(r"\(-?\d+(?:,\d+)?\)\s*–\s*\(?-?\d+(?:,\d+)?\)?|-?\d+(?:,\d+)?\s*–\s*\(?-?\d+(?:,\d+)?\)?|-?\d+(?:,\d+)?|–")


def _fc_num(s: str) -> float:
    s = s.strip("() ")
    ip, _, dp = s.partition(",")
    if len(dp) > 1:              # сноска приклеена к числу: «14,61» при шаге 0,1 → 14,6 (+ сноска 1)
        dp = dp[:1]
    return float(f"{ip}.{dp}" if dp else ip)


def cbr_forecast() -> dict:
    """Среднесрочный прогноз Банка России (последний, PDF с cbr.ru/dkp/mp_dec/decision_key_rate/):
    {asof, url, fact_year, years, rows: {ключ: {'ГГГГ': [мин, макс]}}} — диапазоны по годам
    (инфляция, средняя ключевая, ВВП и компоненты, М2, кредит, платёжный баланс в млрд $)."""
    import pdfplumber
    t = _get(f"{_CBR}/dkp/mp_dec/decision_key_rate/").text
    links = sorted(set(re.findall(r'href="([^"]*forecast_(\d{6})\.pdf)"', t)), key=lambda x: x[1])
    if not links:
        raise RuntimeError("нет ссылки на среднесрочный прогноз ЦБ")
    href, ds = links[-1]
    url = href if href.startswith("http") else _CBR + href
    with pdfplumber.open(io.BytesIO(_get(url).content)) as doc:
        lines = "\n".join((pg.extract_text() or "") for pg in doc.pages).replace("—", "–").split("\n")
    years, fact_year, table, rows = [], None, 0, {}
    for i, line in enumerate(lines):
        s = line.strip()
        if s.startswith("Основные параметры прогноза"):
            table = 1
        elif s.startswith("Показатели платежного баланса"):
            table = 2
        ys = re.findall(r"\b(20\d{2})\b", s)
        if len(ys) >= 3 and re.fullmatch(r"[\d\s]+", s):              # строка лет «2026 2027 2028 2029»
            years = ys
            prev = lines[i - 1].strip() if i else ""
            fact_year = prev if re.fullmatch(r"20\d{2}", prev) else fact_year
            continue
        for key, lbl, tb in _FC_ROWS:
            if key in rows or tb != table:
                continue
            m = re.match(rf"\s*{lbl}[^\d(–]*?(?:\d\s)?(?=\(?-?\d|–\s)", s)
            if not m or not years:
                continue
            rest = s[m.end():]
            toks = _FC_TOK.findall(rest)
            if not toks and i + 1 < len(lines):                       # значения перенесены на след. строку
                toks = _FC_TOK.findall(lines[i + 1])
            n = 1 + len(years)
            while len(toks) > n:                                      # сноски-цифры («Экспорт – 2 …»)
                j = next((k for k, tk in enumerate(toks) if re.fullmatch(r"\d", tk)), None)
                if j is None:
                    break
                toks.pop(j)
            if len(toks) != n:
                continue
            out = {}
            for yy, tk in zip([fact_year] + years, toks):
                if tk == "–" or yy is None:
                    continue
                parts = re.split(r"\s*–\s*", tk) if not tk.startswith("(") else re.findall(r"\(?-?\d+(?:,\d+)?\)?", tk)
                vals = [_fc_num(p) for p in parts if p]
                if vals:
                    out[yy] = [min(vals), max(vals)]
            if out:
                rows[key] = out
    # сноска к ключевой ставке: «с 1 января по 26 июля 2026 года средняя ключевая ставка равна 15,0%,
    # с 27 июля до конца 2026 года … прогнозируется в диапазоне 13,7–14,0%» — прогноз на ОСТАТОК года
    # (сопоставим с текущим уровнем ставки, в отличие от среднего за год)
    flat = re.sub(r"\s+", " ", " ".join(lines))
    rest = None
    m = re.search(r"с (\d{1,2}) ([а-я]+) до конца (\d{4}) года средняя ключевая ставка прогнозируется в диапазоне "
                  r"(\d+(?:,\d+)?)\s*–\s*(\d+(?:,\d+)?)", flat)
    if m and m.group(2) in _MGEN:
        rest = {"year": m.group(3), "from": f"{m.group(3)}-{_MGEN[m.group(2)]:02d}-{int(m.group(1)):02d}",
                "range": [_fc_num(m.group(4)), _fc_num(m.group(5))]}
        ytd = re.search(r"по (\d{1,2}) ([а-я]+) " + m.group(3) + r" года средняя ключевая ставка равна (\d+(?:,\d+)?)", flat)
        if ytd:
            rest["ytd_avg"] = _fc_num(ytd.group(3))
    return {"asof": f"{ds[4:6]}.{ds[2:4]}.20{ds[:2]}", "url": url, "fact_year": fact_year, "years": years, "rows": rows,
            "key_rate_rest": rest}


# Проект федерального бюджета на 2027–2029 годы (млрд руб.). Законопроект № 615-9
# внесён Правительством в Госдуму 30.09.2026. До принятия это НЕ закон: статус и дата
# передаются в API, чтобы интерфейс не выдавал проектные параметры за утверждённый бюджет.
_BUDGET_PLAN = {
    "law": "законопроект № 615-9 о федеральном бюджете на 2027–2029 гг.",
    "status": "проект, внесён в Госдуму 30.09.2026",
    "asof": "2026-09-30",
    "source": "https://sozd.duma.gov.ru/bill/615-9",
    "revenue": {"2027": 43297.2, "2028": 45834.6, "2029": 48620.2},
    "spending": {"2027": 48740.4, "2028": 50954.2, "2029": 53865.4},
    "oilgas": {"2027": 7003.5, "2028": 7633.6, "2029": 7726.1},
    # Публикуем официальные округлённые строки, а не восстанавливаем их разностью:
    # в 2029 году компоненты из пояснительной записки расходятся на 0,1 млрд из-за округления.
    "nonoilgas": {"2027": 36293.7, "2028": 38201.0, "2029": 40894.0},
    "gdp": {"2027": 248323.0, "2028": 263720.0, "2029": 281670.0},
    "deficit_gdp": {"2027": 2.2, "2028": 1.9, "2029": 1.9},
    # Программа внутренних заимствований: размещение ОФЗ, погашение и чистое привлечение.
    "gross_borrowing": {"2027": 7722.0, "2028": 7434.0, "2029": 8004.0},
    "debt_redemption": {"2027": 1323.0, "2028": 1436.0, "2029": 1887.0},
    "net_borrowing": {"2027": 6399.0, "2028": 5998.0, "2029": 6117.0},
}


def budget_plan() -> dict:
    p = _BUDGET_PLAN
    ys = sorted(p["revenue"])
    return {"law": p["law"], "status": p["status"], "asof": p["asof"], "source": p["source"],
            "revenue": _ser(p["revenue"]), "spending": _ser(p["spending"]), "oilgas": _ser(p["oilgas"]),
            "nonoilgas": _ser(p["nonoilgas"]),
            "balance": _ser({y: round(p["revenue"][y] - p["spending"][y], 1) for y in ys}),
            "gdp": _ser(p["gdp"]), "deficit_gdp": _ser(p["deficit_gdp"]),
            "gross_borrowing": _ser(p["gross_borrowing"]),
            "debt_redemption": _ser(p["debt_redemption"]),
            "net_borrowing": _ser(p["net_borrowing"])}


def rs_cpi_weights() -> dict:
    """Веса групп в ИПЦ (структура потребительских расходов, % к итогу) на текущий год — Росстат
    ipc_spr_MM-ГГГГ.xlsx (последний лист). Ключи — названия компонент ЦБ (indicators_cpd), часть
    названий у Росстата другая → алиасы; агрегаты «без …» — разностью."""
    url = _link(f"{_RS}/statistics/price", r"ipc_spr_\d{2}-\d{4}\.xlsx")
    wb = _book(_rosstat_get(url))
    sh = sorted((k for k in wb if re.fullmatch(r"\d{6,7}", k)), key=lambda k: (k[-4:], k[:2]))[-1]
    w = {}
    for r in wb[sh]:
        if len(r) > 2 and isinstance(r[0], str) and isinstance(r[2], (int, float)):
            w.setdefault(re.sub(r"\s+", " ", r[0]).strip().lower(), float(r[2]))
    alias = {"чай, кофе": "чай, кофе, какао", "нефтепродукты": "топливо моторное",
             "жилищно-коммунальные услуги": "жилищные и коммунальные услуги (включая аренду квартир)",
             "продовольствие без плодоовощей": "продовольственные товары (без овощей, картофеля и фруктов)",
             "базовый ипц": "базовый индекс потребительских цен"}
    get = lambda nm: w.get(alias.get(nm.lower(), nm.lower()))  # noqa: E731
    out = {}
    names = ["Все товары и услуги", "Продовольственные товары", "Продовольствие без плодоовощей", "Мясопродукты", "Рыбопродукты",
             "Масло и жиры", "Молоко и молочная продукция", "Сыр", "Яйца", "Сахар", "Кондитерские изделия", "Чай, кофе",
             "Хлеб и хлебобулочные изделия", "Макаронные и крупяные изделия", "Плодоовощная продукция, включая картофель",
             "Алкогольные напитки", "Общественное питание", "Непродовольственные товары", "Одежда и белье", "Меха и меховые изделия",
             "Трикотажные изделия", "Обувь кожаная, текстильная и комбинированная", "Моющие и чистящие средства",
             "Парфюмерно-косметические товары", "Галантерея", "Табачные изделия", "Мебель", "Электротовары и другие бытовые приборы",
             "Печатные издания", "Телерадиотовары", "Персональные компьютеры", "Средства связи", "Строительные материалы",
             "Легковые автомобили", "Инструменты и оборудование", "Нефтепродукты", "Медицинские товары", "Услуги", "Бытовые услуги",
             "Услуги пассажирского транспорта", "Жилищно-коммунальные услуги", "Услуги гостиниц и прочих мест проживания",
             "Услуги в системе образования", "Услуги организаций культуры", "Услуги в сфере зарубежного туризма",
             "Экскурсионные услуги", "Санаторно-оздоровительные услуги", "Медицинские услуги", "Базовый ИПЦ"]
    for nm in names:
        v = get(nm)
        if v is not None:
            out[nm] = round(v, 3)
    if "Непродовольственные товары" in out and "Нефтепродукты" in out:
        out["Непрод. без нефтепродуктов"] = round(out["Непродовольственные товары"] - out["Нефтепродукты"], 3)
    if "Услуги" in out and "Жилищно-коммунальные услуги" in out:
        out["Услуги без ЖКХ"] = round(out["Услуги"] - out["Жилищно-коммунальные услуги"], 3)
    # услуги связи: у Росстата группы «Услуги связи» нет — телекоммуникационные + почтовые
    tel = sum(v for k, v in w.items() if k in ("услуги телекоммуникационные", "услуги почтовой связи"))
    if tel:
        out["Услуги связи"] = round(tel, 3)
    # базовый ИПЦ, м/м (индекс к пред. месяцу) по всем месячным листам — для г/г за месяц, которого
    # ещё нет в PDF ЦБ «Инфляция в России» (Росстат публикует раньше)
    core_mom = {}
    for sn, rows in wb.items():
        if re.fullmatch(r"\d{6,7}", sn):
            for r in rows:
                if r and isinstance(r[0], str) and r[0].strip().lower().startswith("базовый индекс потребительских цен") and len(r) > 3 and isinstance(r[3], (int, float)):
                    core_mom[f"{sn[-4:]}-{int(sn[:2]):02d}"] = r[3]
                    break
    return {"year": sh[-4:], "src": url, "w": out, "core_mom": dict(sorted(core_mom.items()))}


def moex_cap_years(y0: int = 2013) -> list:
    """Капитализация рынка акций Мосбиржи на последний торговый день каждого года, млрд ₽ —
    для индикатора Баффета за годы, которых ещё нет у World Bank. Выбросы ISS (2020–2021 —
    ошибка единиц, ×30) отсекаем: скачок >3× к соседям."""
    from concurrent.futures import ThreadPoolExecutor
    import json
    import os
    # прошлые годы не меняются → дисковый кэш; сеть только для лет, которых в кэше нет
    # (раньше каждое обновление макро делало до 78 запросов к ISS)
    path = os.path.join(os.path.dirname(__file__), "data", "moex_cap_years.json")
    try:
        with open(path, encoding="utf-8") as f:
            cached = {r["date"]: r for r in json.load(f)}
    except Exception:  # noqa: BLE001
        cached = {}

    def one(y):
        if str(y) in cached:
            return cached[str(y)]
        for dd in (31, 30, 29, 28, 27, 26):
            t = _get(f"https://iss.moex.com/iss/statistics/engines/stock/capitalization.csv?date={y}-12-{dd}").text
            m = re.search(r"TRADEDATE\s*\n([\d.]+);([\d-]+)", t)
            if m and m.group(2).startswith(str(y)):
                return {"date": str(y), "val": round(float(m.group(1)) / 1e9, 1), "tradedate": m.group(2)}
        return None
    with ThreadPoolExecutor(max_workers=6) as ex:
        rows = [r for r in ex.map(one, range(y0, datetime.now().year)) if r]
    if rows and len(rows) > len(cached):
        try:
            _dump_json(path, rows)
        except Exception as e:  # noqa: BLE001
            log.warning("moex_cap_years cache: %s", e)
    med = sorted(r["val"] for r in rows)[len(rows) // 2] if rows else 0
    return [r for r in rows if r["val"] < med * 3]


EXT_SOURCES = [
    ("cpi_weights", rs_cpi_weights, {}), ("moex_cap_years", moex_cap_years, []),
    ("cbr_forecast", cbr_forecast, {}), ("budget_plan", budget_plan, {}), ("weekly_items", rosstat_weekly_items, {}),
    ("core_infl", cbr_core_inflation, {}), ("gov_debt", minfin_gov_debt, {}),
    ("cpi_comp", rs_cpi_components, {}), ("ppi", rs_ppi, {}), ("industrial", rs_industrial, {}),
    ("ivbo", rs_ivbo, []), ("retail", rs_retail, {}), ("wages", rs_wages, {}),
    ("real_income", rs_real_income, {}), ("unemployment", rs_unemployment, {}),
    ("housing", rs_housing, {}), ("confidence", rs_confidence, {}), ("gdp_use", rs_gdp_use, {}),
    ("reserves", cbr_reserves, {}), ("keyrate_hist", cbr_keyrate_hist, []),
    ("liquidity", cbr_liquidity, []), ("money_base", cbr_money_base, {}),
    ("credit_econ", cbr_credit_economy, {}), ("households", cbr_households, {}),
    ("nfs", cbr_nfs, {}), ("ext_debt", cbr_ext_debt, {}), ("trade", cbr_trade, {}),
    ("rate_terms", cbr_rate_terms, {}), ("issuance", cbr_issuance, {}),
]
