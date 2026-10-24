"""Результаты аукционов ОФЗ Минфина: дата, выпуск, тип, объём предложения, спрос,
цена отсечения, средневзвешенная цена/доходность.

Источник — годовые xlsx Минфина (minfin.gov.ru, обновляются раз в неделю). Ссылки на
файлы берём со страницы аукционов (в имени файла — год), парсим openpyxl, кэшируем."""
import io
import logging
import re
import time
from datetime import datetime

import requests

log = logging.getLogger(__name__)

_UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15"
_PAGE = "https://minfin.gov.ru/ru/perfomance/public_debt/internal/operations/ofz/auction"
_BASE = "https://minfin.gov.ru"
_TTL = 6 * 3600
_cache: dict = {"at": 0.0, "data": None}      # (устар.) общий кэш; теперь — по числу лет, см. _caches
_caches: dict = {}                             # years → {"at", "data", "asof"}


def _xlsx_links() -> list[str]:
    r = requests.get(_PAGE, headers={"User-Agent": _UA}, timeout=30)
    r.raise_for_status()
    links = re.findall(r'href="(/common/upload/[^"]*Auction_Results[^"]*\.xlsx)"', r.text)
    # уникальные, порядок сохраняем
    seen, out = set(), []
    for l in links:
        if l not in seen:
            seen.add(l)
            out.append(_BASE + l)
    return out


def _parse_xlsx(content: bytes) -> list[dict]:
    """Раскладка столбцов у Минфина менялась (в свежих файлах есть колонка «Формат»,
    в старых нет) — поэтому находим строку заголовка и мапим столбцы ПО НАЗВАНИЯМ."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    hdr_i = None
    for i, r in enumerate(rows):
        joined = " ".join(str(c or "").lower() for c in r)
        if "код" in joined and "выпуска" in joined and "отсечен" in joined:
            hdr_i = i
            break
    if hdr_i is None:
        return []
    hdr = [re.sub(r"\s+", " ", str(c or "").lower()).strip() for c in rows[hdr_i]]

    def col(*subs, exclude=None):
        for j, h in enumerate(hdr):
            if all(sub in h for sub in subs) and not (exclude and exclude in h):
                return j
        return None

    ci = {
        "date": col("дата", exclude="погашен"),        # дата аукциона (не погашения)
        "code": col("код", "выпуска"),
        "type": col("тип"),
        "matdate": col("дата", "погашен"),
        "offered": col("объем", "предложен"),
        "cutoff_price": col("цена", "отсечен"),
        "wa_price": col("цена", "средневзв"),
        "cutoff_yield": col("доходность", "отсеч"),
        "wa_yield": col("доходность", "средневзв"),
        "demand": col("объем", "спрос"),
        "placed": col("объем", "размещен"),        # фактически размещено по номиналу
        "proceeds": col("объем", "выручк"),         # выручка (денег получено), млн ₽
        "format": col("формат"),                    # «Аукцион» / «ДРПА» (доразмещение после аукциона)
    }
    if ci["date"] is None or ci["code"] is None:
        return []

    def cell(r, key):
        j = ci.get(key)
        return r[j] if (j is not None and j < len(r)) else None

    out = []
    for r in rows[hdr_i + 1:]:
        d = cell(r, "date")
        code = cell(r, "code")
        if not (hasattr(d, "year") and code and "RMFS" in str(code)):
            continue

        def f(key):
            v = cell(r, key)
            try:
                return round(float(v), 4) if v not in (None, "") else None
            except (TypeError, ValueError):
                return None

        cd = str(code).strip()
        mat = cell(r, "matdate")
        typ = str(cell(r, "type") or "").strip()
        demand, offered = f("demand"), f("offered")
        wa_price, cutoff_price = f("wa_price"), f("cutoff_price")
        wa_yield = f("wa_yield")
        # ОФЗ-ПК (флоатеры) — доходности нет; считаем simple margin из цены размещения
        # (у ОФЗ-ПК квотируемая надбавка ≈0, поэтому SM ≈ pull-to-par по формуле Фабоцци)
        sm_bp = None
        if "ПК" in typ and wa_price and mat and hasattr(mat, "year"):
            try:
                yrs = (mat.date() - d.date()).days / 365.0
                if yrs > 0 and wa_price > 0:
                    sm_bp = round((100.0 - wa_price) / yrs * (100.0 / wa_price) * 100.0)
            except (AttributeError, ValueError, ZeroDivisionError):
                sm_bp = None
        fmt_ = str(cell(r, "format") or "Аукцион").strip()
        out.append({
            "date": d.strftime("%Y-%m-%d"),
            "format": fmt_,
            "secid": "SU" + cd,
            "series": cd.replace("RMFS", ""),
            "type": typ,
            "matdate": mat.strftime("%Y-%m-%d") if hasattr(mat, "year") else None,
            "offered_mln": offered,
            "demand_mln": demand,
            "placed_mln": f("placed"),          # фактически размещено (номинал)
            "proceeds_mln": f("proceeds"),      # выручка, млн ₽ (у старых файлов может не быть)
            "bid_to_cover": round(demand / offered, 2) if (demand and offered) else None,
            "cutoff_price": cutoff_price,
            "wa_price": wa_price,
            "cutoff_yield": f("cutoff_yield"),
            "wa_yield": wa_yield,
            "sm_bp": sm_bp,
        })
    return out


def get_auctions(force: bool = False, years: int = 2) -> list[dict]:
    """Список аукционов за последние `years` лет, новые сверху. Кэш на 6ч."""
    now = time.time()
    # кэш — отдельно на каждый период: вкладка аукционов берёт 2 года, макро — с 2021 года
    cc = _caches.setdefault(years, {"at": 0.0, "data": None, "asof": None})
    if not force and cc["data"] is not None and now - cc["at"] < _TTL:
        return cc["data"]
    try:
        links = _xlsx_links()
    except Exception as e:  # noqa: BLE001
        log.warning("minfin auction page failed: %s", e)
        return cc["data"] or []

    cur = datetime.today().year
    want_years = list(range(cur - years + 1, cur + 1))
    # фильтруем по году ИМЕННО в имени файла (…_rus_2026_…), не по пути
    urls = [l for l in links if any(f"_rus_{y}_" in l for y in want_years)]

    recs: list[dict] = []
    for url in urls:
        try:
            c = requests.get(url, headers={"User-Agent": _UA}, timeout=40).content
            recs += _parse_xlsx(c)
        except Exception as e:  # noqa: BLE001
            log.warning("auction xlsx %s failed: %s", url, e)

    # ключ — с форматом: в один день у выпуска бывают и аукцион, и ДРПА (доразмещение); раньше
    # ключ был «дата + выпуск», и одна из двух строк затиралась — объёмы размещения занижались
    seen = {}
    for r in recs:
        seen[(r["date"], r["secid"], r.get("format") or "")] = r
    data = sorted(seen.values(), key=lambda x: (x["date"], x["secid"]), reverse=True)
    if data:
        # по состоянию на — дата в имени самого свежего файла (…_rus_2026_20260924.xlsx)
        dts = [m.group(1) for u in urls for m in [re.search(r"_(\d{8})\.xlsx", u)] if m]
        asof = max(dts) if dts else None
        cc.update(at=now, data=data, asof=f"{asof[:4]}-{asof[4:6]}-{asof[6:]}" if asof else None)
    return data


def asof(years: int = 2) -> str | None:
    """Дата, по состоянию на которую Минфин выложил результаты (файл обновляется раз в неделю)."""
    return (_caches.get(years) or {}).get("asof")
